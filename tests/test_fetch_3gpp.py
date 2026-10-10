import io
import json
import zipfile

import httpx
import pytest

pytest.importorskip("docx")
from docx import Document

from app.chunking import chunk_markdown
from scripts import fetch_3gpp as f


def make_docx() -> bytes:
    d = Document()
    d.add_heading("1 Scope", level=1)
    d.add_paragraph("The present document defines KPIs.")
    d.add_heading("5.1.1.6 Handover success rate", level=3)
    d.add_paragraph("HO success rate = successful / attempted.")
    d.add_heading("5.1.1.6.1 Detail", level=4)
    d.add_paragraph("")  # empty paragraphs must vanish
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Counter", "Meaning"
    t.cell(1, 0).text, t.cell(1, 1).text = "HO.Att", "attempts | total"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_version_decoding():
    assert f.version_from_zip("28554-ia0.zip") == "18.10.0"
    assert f.version_from_zip("38331-ia30.zip".replace("a30", "a3")) == "18.10.3"
    assert f.version_from_zip("32450-i00.zip") == "18.0.0"


def test_find_zip_picks_the_spec_not_its_neighbours():
    names = ["28552-ia0.zip", "28554-i90.zip", "28554-ia0.zip", "285540-i00.zip"]
    assert f.find_zip(names, "28.554") == "28554-ia0.zip"
    assert f.find_zip(names, "28.999") is None


def test_docx_to_markdown_structure_and_table():
    md = f.docx_to_markdown(make_docx(), "28.554", "18.10.0", "Rel-18", "http://x")
    assert md.startswith("# 3GPP TS 28.554 Rel-18 v18.10.0 (http://x)")
    assert "\n## 1 Scope" in md and "\n## 5.1.1.6 Handover success rate" in md
    assert "**5.1.1.6.1 Detail**" in md and "\n## 5.1.1.6.1" not in md  # deep clause: no split
    assert "| Counter | Meaning |" in md and "|---|---|" in md
    assert "attempts / total" in md  # pipe inside a cell must not break the table
    chunks = chunk_markdown(md)
    assert len(chunks) == 2 and all(c.startswith("# 3GPP TS 28.554") for c in chunks)


def zip_with(name: str, data: bytes) -> bytes:
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr(name, data)
    return b.getvalue()


def client_for(zip_bytes: bytes, listing: str) -> httpx.Client:
    def handler(req: httpx.Request) -> httpx.Response:
        if req.url.path.endswith("/"):
            return httpx.Response(200, text=listing)
        return httpx.Response(200, content=zip_bytes)
    return httpx.Client(transport=httpx.MockTransport(handler))


def test_fetch_spec_writes_markdown(tmp_path, monkeypatch):
    monkeypatch.setattr(f, "docx_to_markdown", lambda *a: "# t\n" + "x" * 30_000)
    c = client_for(zip_with("28554-ia0.docx", make_docx()),
                   '<a href="/ftp/Specs/latest/Rel-18/28_series/28554-ia0.zip">x</a>')
    res = f.fetch_spec(c, "28.554", "Rel-18", tmp_path)
    assert res["status"] == "ok" and res["version"] == "18.10.0"
    assert (tmp_path / "TS_28_554_Rel-18.md").read_text().startswith("# t")
    assert len(res["sha256"]) == 64


def test_legacy_doc_only_is_skipped_not_garbled(tmp_path, monkeypatch):
    monkeypatch.setattr(f, "_doc_to_docx", lambda raw: None)  # no LibreOffice
    c = client_for(zip_with("28554-ia0.doc", b"\xd0\xcf\x11\xe0"),
                   '<a href="28554-ia0.zip">x</a>')
    assert f.fetch_spec(c, "28.554", "Rel-18", tmp_path)["status"] == "no_docx"
    assert not list(tmp_path.glob("*.md"))


def test_spec_missing_from_listing(tmp_path):
    c = client_for(b"", '<a href="28552-ia0.zip">x</a>')
    assert f.fetch_spec(c, "28.554", "Rel-18", tmp_path)["status"] == "not_found"


def test_main_writes_manifest_and_fails_when_nothing_converted(tmp_path, monkeypatch):
    monkeypatch.setattr(f, "fetch_spec", lambda *a, **k: {"spec": "28.554", "status": "not_found"})
    assert f.main(["--specs", "28.554", "--out", str(tmp_path)]) == 1
    m = json.loads((tmp_path / "MANIFEST.json").read_text())
    assert m["specs"][0]["status"] == "not_found"


def test_body_inside_content_control_is_not_lost():
    """Some specs wrap the whole body in a w:sdt; a direct-children walk saw it as empty
    (TS 38.133 converted to ~2.6k chars in CI)."""
    from docx.oxml import parse_xml
    d = Document()
    body = d.element.body
    sdt = parse_xml(
        '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:sdtContent><w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
        '<w:r><w:t>4 Hidden clause</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>Body text inside the control.</w:t></w:r></w:p></w:sdtContent></w:sdt>')
    body.insert(0, sdt)
    buf = io.BytesIO()
    d.save(buf)
    md = f.docx_to_markdown(buf.getvalue(), "38.133", "18.15.0", "Rel-18", "http://x")
    assert "## 4 Hidden clause" in md and "Body text inside the control." in md


def test_suspiciously_small_conversion_is_rejected(tmp_path):
    c = client_for(zip_with("28554-ia0.docx", make_docx()), '<a href="28554-ia0.zip">x</a>')
    res = f.fetch_spec(c, "28.554", "Rel-18", tmp_path)
    assert res["status"] == "too_small" and not list(tmp_path.glob("*.md"))


def test_legacy_doc_converts_end_to_end_with_libreoffice(tmp_path):
    """Real LibreOffice round trip: .docx -> .doc (as old 3GPP zips ship) -> _doc_to_docx."""
    import shutil
    import subprocess
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed")
    src = tmp_path / "spec.docx"
    src.write_bytes(make_docx())
    subprocess.run(["soffice", "--headless", "--convert-to", "doc", "--outdir", str(tmp_path),
                    str(src)], check=True, capture_output=True, timeout=300)
    data = f._doc_to_docx((tmp_path / "spec.doc").read_bytes())
    assert data
    md = f.docx_to_markdown(data, "32.450", "18.0.0", "Rel-18", "http://x")
    assert "## 1 Scope" in md and "HO success rate" in md
def test_text_fallback_turns_clause_headings_into_chunk_boundaries():
    txt = "Foreword\n\n1\tScope\nThis spec defines KPIs.\n5.1.1 Cell availability\n5.1.1.1 Deeper clause\nBody text here.\n" \
          "Table 1 shows values.\n6 Annex A.\n"
    md = f.text_to_markdown(txt, "32.450", "18.0.0", "Rel-18", "u")
    assert "## 1\tScope" in md
    assert "## 5.1.1 Cell availability" in md
    assert "## 5.1.1.1" not in md  # level 4 stays inside its parent, as in the docx path
    assert md.startswith("# 3GPP TS 32.450 Rel-18 v18.0.0")
    assert "\nTable 1 shows values." in md  # not mistaken for a clause

